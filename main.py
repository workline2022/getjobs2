#!/usr/bin/env python3
"""
Job Sourcing Platform CLI

Usage:
    python main.py scrape [--layer1-only] [--layer2-only] [--query QUERY]
    python main.py query  [--state CA] [--industry tech] [--company NAME] [--limit 50]
    python main.py export [--state CA] [--industry tech] [--company NAME] [--out jobs.xlsx]
    python main.py stats
    python main.py clear

Examples:
    python main.py scrape
    python main.py scrape --layer1-only
    python main.py query --state CA --industry tech
    python main.py export --state NY --industry healthcare --out ny_healthcare.xlsx
    python main.py export --out all_jobs.xlsx
    python main.py stats
"""

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

import database
import layer1
import layer2
from filters import filter_by_state, filter_by_industry
from models import Job

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(name)s — %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("main")


def cmd_scrape(args: argparse.Namespace) -> None:
    database.init_db()

    all_jobs: list[Job] = []
    layer1_keys: set[str] = set()

    if not args.layer2_only:
        print("=== Layer 1: ATS Direct Scraping ===")
        l1_jobs = layer1.scrape_all()
        all_jobs.extend(l1_jobs)
        layer1_keys = {j.dedup_key() for j in l1_jobs}
        print(f"Layer 1 complete: {len(l1_jobs)} unique jobs")

    if not args.layer1_only:
        query = getattr(args, "query", None) or "engineer"
        location = getattr(args, "location", None) or "United States"
        print(f"\n=== Layer 2: JobSpy (query='{query}', location='{location}') ===")
        l2_jobs = layer2.scrape_jobspy(
            query=query,
            location=location,
            results_wanted=50,
            existing_keys=layer1_keys,
        )
        all_jobs.extend(l2_jobs)
        print(f"Layer 2 complete: {len(l2_jobs)} new jobs after dedup")

    if all_jobs:
        inserted, skipped = database.upsert_jobs(all_jobs)
        print(f"\nDatabase: {inserted} inserted, {skipped} skipped (duplicates)")
    else:
        print("No jobs to save.")

    print(f"Total jobs in DB: {database.count_jobs()}")


def cmd_query(args: argparse.Namespace) -> None:
    database.init_db()

    state = (args.state or "").upper()
    industry = (args.industry or "").lower()
    company = args.company or ""
    limit = args.limit

    rows = database.query_jobs(
        state=state,
        industry=industry,
        company=company,
        limit=limit,
    )

    if not rows:
        print("No jobs found matching the given filters.")
        return

    # Pretty print
    print(f"\nFound {len(rows)} job(s):\n")
    for i, row in enumerate(rows, 1):
        print(f"[{i}] {row['title']}")
        print(f"     Company  : {row['company']}")
        print(f"     Location : {row['location']}  (state: {row['state'] or 'n/a'})")
        print(f"     Industry : {row['industry']}")
        print(f"     Source   : {row['source']}")
        print(f"     URL      : {row['url']}")
        print()


def cmd_stats(args: argparse.Namespace) -> None:
    database.init_db()
    total = database.count_jobs()
    print(f"Total jobs in DB: {total}")

    import sqlite3
    from database import get_connection
    conn = get_connection()

    print("\nBy source:")
    for row in conn.execute("SELECT source, COUNT(*) AS n FROM jobs GROUP BY source ORDER BY n DESC").fetchall():
        print(f"  {row['source']:<15} {row['n']}")

    print("\nBy industry:")
    for row in conn.execute("SELECT industry, COUNT(*) AS n FROM jobs GROUP BY industry ORDER BY n DESC").fetchall():
        print(f"  {row['industry'] or 'other':<15} {row['n']}")

    print("\nTop states:")
    for row in conn.execute(
        "SELECT state, COUNT(*) AS n FROM jobs WHERE state != '' GROUP BY state ORDER BY n DESC LIMIT 10"
    ).fetchall():
        print(f"  {row['state']:<6} {row['n']}")

    conn.close()


def cmd_export(args: argparse.Namespace) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("ERROR: openpyxl is required. Run: pip install openpyxl")
        sys.exit(1)

    database.init_db()

    state = (args.state or "").upper()
    industry = (args.industry or "").lower()
    company = args.company or ""
    limit = args.limit

    rows = database.query_jobs(
        state=state,
        industry=industry,
        company=company,
        limit=limit,
    )

    if not rows:
        print("No jobs found matching the given filters.")
        return

    # Determine output path
    out_path = Path(args.out)

    # Build workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Jobs"

    # Header style
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    columns = [
        ("Title",       "title",       45),
        ("Company",     "company",     22),
        ("Location",    "location",    28),
        ("State",       "state",        7),
        ("Industry",    "industry",    14),
        ("Department",  "department",  22),
        ("Source",      "source",      13),
        ("URL",         "url",         55),
        ("Scraped At",  "inserted_at", 18),
    ]

    # Write headers
    for col_idx, (header, _, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    # Alternating row fills
    fill_even = PatternFill("solid", fgColor="DCE6F1")
    url_font = Font(color="1155CC", underline="single")
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(vertical="center", wrap_text=True)

    for row_idx, row in enumerate(rows, start=2):
        fill = fill_even if row_idx % 2 == 0 else None
        for col_idx, (_, field, _) in enumerate(columns, start=1):
            value = row.get(field, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if fill:
                cell.fill = fill
            # Hyperlink for URL column
            if field == "url" and value:
                cell.hyperlink = value
                cell.font = url_font
                cell.alignment = wrap
            elif field in ("state", "source"):
                cell.alignment = center
            else:
                cell.alignment = wrap
        ws.row_dimensions[row_idx].height = 18

    # Auto-filter on header row
    ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 14

    summary_header_font = Font(bold=True, color="FFFFFF")
    summary_fill = PatternFill("solid", fgColor="1F4E79")

    def _write_summary_section(title, data_rows, start_row):
        ws2.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=12)
        start_row += 1
        for label, val_col in [("Category", "Count")]:
            c1 = ws2.cell(row=start_row, column=1, value=label)
            c2 = ws2.cell(row=start_row, column=2, value=val_col)
            for c in (c1, c2):
                c.font = summary_header_font
                c.fill = summary_fill
                c.alignment = Alignment(horizontal="center")
        start_row += 1
        for k, v in data_rows:
            ws2.cell(row=start_row, column=1, value=k)
            ws2.cell(row=start_row, column=2, value=v).alignment = Alignment(horizontal="center")
            start_row += 1
        return start_row + 1

    from collections import Counter
    by_industry = Counter(r["industry"] or "other" for r in rows)
    by_state    = Counter(r["state"] or "n/a"    for r in rows)
    by_source   = Counter(r["source"]             for r in rows)
    by_company  = Counter(r["company"]            for r in rows)

    row_cursor = 1
    ws2.cell(row=row_cursor, column=1, value=f"Export generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    ws2.cell(row=row_cursor, column=1).font = Font(italic=True)
    row_cursor += 1
    ws2.cell(row=row_cursor, column=1, value=f"Total jobs: {len(rows)}")
    ws2.cell(row=row_cursor, column=1).font = Font(bold=True)
    row_cursor += 2

    row_cursor = _write_summary_section("By Industry", by_industry.most_common(), row_cursor)
    row_cursor = _write_summary_section("By State (top 15)", by_state.most_common(15), row_cursor)
    row_cursor = _write_summary_section("By Source", by_source.most_common(), row_cursor)
    row_cursor = _write_summary_section("By Company (top 20)", by_company.most_common(20), row_cursor)

    wb.save(out_path)
    print(f"Exported {len(rows)} jobs to: {out_path.resolve()}")
    print(f"  Sheets: 'Jobs' (data) + 'Summary' (breakdown by industry/state/source/company)")


def cmd_clear(args: argparse.Namespace) -> None:
    n = database.delete_all_jobs()
    print(f"Deleted {n} job(s) from the database.")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Job Sourcing Platform",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    sub = parser.add_subparsers(dest="command", required=True)

    # scrape
    p_scrape = sub.add_parser("scrape", help="Scrape jobs from ATS + aggregators")
    p_scrape.add_argument("--layer1-only", action="store_true", help="Only run Layer 1 (ATS direct)")
    p_scrape.add_argument("--layer2-only", action="store_true", help="Only run Layer 2 (JobSpy)")
    p_scrape.add_argument("--query", default="engineer", help="Search query for JobSpy (Layer 2)")
    p_scrape.add_argument("--location", default="United States", help="Location for JobSpy (Layer 2)")

    # query
    p_query = sub.add_parser("query", help="Query stored jobs")
    p_query.add_argument("--state", help="Filter by US state code (e.g. CA, NY)")
    p_query.add_argument("--industry", help="Filter by industry (tech, healthcare, finance, etc.)")
    p_query.add_argument("--company", help="Filter by company name (partial match)")
    p_query.add_argument("--limit", type=int, default=50, help="Max results to display")

    # export
    p_export = sub.add_parser("export", help="Export jobs to an Excel (.xlsx) file")
    p_export.add_argument("--state", help="Filter by US state code (e.g. CA, NY)")
    p_export.add_argument("--industry", help="Filter by industry (tech, healthcare, finance, etc.)")
    p_export.add_argument("--company", help="Filter by company name (partial match)")
    p_export.add_argument("--limit", type=int, default=10000, help="Max rows to export (default: 10000)")
    p_export.add_argument("--out", default="jobs.xlsx", help="Output file path (default: jobs.xlsx)")

    # stats
    sub.add_parser("stats", help="Show database statistics")

    # clear
    sub.add_parser("clear", help="Delete all jobs from the database")

    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    dispatch = {
        "scrape": cmd_scrape,
        "query": cmd_query,
        "export": cmd_export,
        "stats": cmd_stats,
        "clear": cmd_clear,
    }
    dispatch[args.command](args)


if __name__ == "__main__":
    main()
